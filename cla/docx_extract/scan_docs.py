# -*- coding: utf-8 -*-
import zipfile, os, io
import docx

files = [r'c:\Users\admin\Downloads\文档\付志成劳动实践报告.docx',
         r'c:\Users\admin\Downloads\文档\何瑞淇劳动实践报告.docx',
         r'c:\Users\admin\Downloads\文档\王若涵.docx']

def repair(src):
    """重新打包，绕过坏CRC"""
    buf = io.BytesIO()
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            with zin.open(item) as f:
                f._expected_crc = None
                data = f.read()
            zout.writestr(item, data)
    buf.seek(0)
    return buf

for f in files:
    try:
        d = docx.Document(f)
        src = f
    except Exception:
        d = docx.Document(repair(f))
        src = '(repaired)'
    print('=' * 70)
    print(os.path.basename(f), src, '| paras:', len(d.paragraphs), '| tables:', len(d.tables))
    for i, p in enumerate(d.paragraphs):
        if p.text.strip():
            print('  P%d:' % i, p.text.strip()[:90])
    for ti, tb in enumerate(d.tables):
        print('  TABLE%d rows=%d cols=%d' % (ti, len(tb.rows), len(tb.columns)))
        for ri, row in enumerate(tb.rows):
            for ci, c in enumerate(row.cells):
                txt = c.text.strip().replace('\n', ' / ')
                has_img = 'blip' in c._tc.xml or 'graphic' in c._tc.xml
                print('    r%dc%d img=%s | %s' % (ri, ci, has_img, txt[:70]))
